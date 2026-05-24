"""
utils/file_download.py — Artifact retrieval and metadata extraction.

Supported metadata extractors:
  - PDF   → ``pypdf`` (falls back to ``PyPDF2`` if ``pypdf`` is absent)
  - DOCX  → ``python-docx`` core properties
  - XLSX  → ``openpyxl`` workbook properties
  - Images → ``exifread`` EXIF
"""

import os
import datetime
import hashlib
import requests
from urllib.parse import urlparse
from rich.table import Table
import exifread

from cli.ui import console, THEME, print_success, print_error, print_warn, print_info
from core.config import DIR_DOWNLOADS


# ---------------------------------------------------------------------------
# PDF backend — prefer ``pypdf``, fall back to ``PyPDF2``
# ---------------------------------------------------------------------------

try:
    from pypdf import PdfReader as _PdfReader  # modern replacement
    _PDF_BACKEND = "pypdf"
except ImportError:
    try:
        from PyPDF2 import PdfReader as _PdfReader  # legacy
        _PDF_BACKEND = "PyPDF2"
    except ImportError:
        _PdfReader = None
        _PDF_BACKEND = None


# ---------------------------------------------------------------------------
# Optional metadata backends
# ---------------------------------------------------------------------------

try:
    from docx import Document as _DocxDocument  # python-docx
    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False

try:
    import openpyxl  # already in requirements for Excel export
    _HAS_OPENPYXL = True
except ImportError:
    _HAS_OPENPYXL = False


# ---------------------------------------------------------------------------
# FileDownload service
# ---------------------------------------------------------------------------

class FileDownload:
    """
    Service layer for artifact retrieval and metadata extraction.
    """

    def __init__(self, destination_dir: str = DIR_DOWNLOADS):
        self.directory = destination_dir
        os.makedirs(self.directory, exist_ok=True)

    # ---------------------- Metadata extractors ----------------------------

    def _render_table(self, title: str, rows: list[tuple[str, str]]) -> None:
        if not rows:
            console.print(f"  [{THEME['WARN']}]⚠[/]  No metadata found ({title}).")
            return
        tbl = Table(title=title, show_header=False, box=None)
        tbl.add_column("Field", style=THEME["PRIMARY"])
        tbl.add_column("Value", style="green")
        for k, v in rows:
            tbl.add_row(str(k), str(v))
        console.print(tbl)

    def _extract_pdf_metadata(self, file_path: str) -> None:
        if _PdfReader is None:
            print_warn("No PDF backend installed (pypdf / PyPDF2).")
            return
        try:
            reader = _PdfReader(file_path)
            meta = reader.metadata or {}
            pages = len(reader.pages)

            rows: list[tuple[str, str]] = [("Pages", str(pages))]
            meta_map = {
                "/Author": "Author",
                "/Creator": "Creator (Software)",
                "/Producer": "Producer (Software)",
                "/Subject": "Subject",
                "/Title": "Title",
                "/Keywords": "Keywords",
                "/CreationDate": "Creation Date",
                "/ModDate": "Modification Date",
            }
            # meta supports both dict and attribute access depending on version
            for key, label in meta_map.items():
                value = None
                try:
                    value = meta.get(key) if hasattr(meta, "get") else None
                except Exception:
                    value = None
                if value:
                    rows.append((label, str(value)))

            self._render_table(
                f"PDF Metadata — {os.path.basename(file_path)} (via {_PDF_BACKEND})",
                rows,
            )
        except Exception as e:
            print_error(f"PDF metadata error: {e}")

    def _extract_exif_metadata(self, file_path: str) -> None:
        try:
            with open(file_path, "rb") as f:
                tags = exifread.process_file(f, details=False)
            skip = {"JPEGThumbnail", "TIFFThumbnail", "Filename", "EXIF MakerNote"}
            rows = [(tag, str(value)) for tag, value in tags.items() if tag not in skip]
            self._render_table(f"EXIF — {os.path.basename(file_path)}", rows)
        except Exception as e:
            print_error(f"EXIF metadata error: {e}")

    def _extract_docx_metadata(self, file_path: str) -> None:
        if not _HAS_DOCX:
            print_warn("python-docx not installed — cannot read DOCX metadata.")
            return
        try:
            doc = _DocxDocument(file_path)
            props = doc.core_properties
            rows = [
                ("Title",            props.title or ""),
                ("Author",           props.author or ""),
                ("Last modified by", props.last_modified_by or ""),
                ("Created",          str(props.created) if props.created else ""),
                ("Modified",         str(props.modified) if props.modified else ""),
                ("Subject",          props.subject or ""),
                ("Keywords",         props.keywords or ""),
                ("Category",         props.category or ""),
                ("Comments",         props.comments or ""),
                ("Revision",         str(props.revision) if props.revision else ""),
                ("Paragraphs",       str(len(doc.paragraphs))),
            ]
            rows = [(k, v) for k, v in rows if v]
            self._render_table(f"DOCX — {os.path.basename(file_path)}", rows)
        except Exception as e:
            print_error(f"DOCX metadata error: {e}")

    def _extract_xlsx_metadata(self, file_path: str) -> None:
        if not _HAS_OPENPYXL:
            print_warn("openpyxl not installed — cannot read XLSX metadata.")
            return
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            props = wb.properties
            rows = [
                ("Title",    props.title or ""),
                ("Author",   props.creator or ""),
                ("Last modified by", props.lastModifiedBy or ""),
                ("Created",  str(props.created) if props.created else ""),
                ("Modified", str(props.modified) if props.modified else ""),
                ("Subject",  props.subject or ""),
                ("Keywords", props.keywords or ""),
                ("Description", props.description or ""),
                ("Category", props.category or ""),
                ("Company",  getattr(props, "company", "") or ""),
                ("Sheets",   ", ".join(wb.sheetnames)),
            ]
            rows = [(k, v) for k, v in rows if v]
            self._render_table(f"XLSX — {os.path.basename(file_path)}", rows)
        except Exception as e:
            print_error(f"XLSX metadata error: {e}")

    def _extract_basic_metadata(self, file_path: str) -> None:
        try:
            stat = os.stat(file_path)
            size_kb = stat.st_size / 1024
            mtime = datetime.datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds")
            # Hash first 1 MB for quick fingerprint
            sha = hashlib.sha256()
            with open(file_path, "rb") as f:
                sha.update(f.read(1024 * 1024))
            rows = [
                ("Size",         f"{size_kb:,.1f} KB ({stat.st_size:,} bytes)"),
                ("Modified",     mtime),
                ("SHA-256 (1M)", sha.hexdigest()),
            ]
            self._render_table(f"Filesystem — {os.path.basename(file_path)}", rows)
        except Exception as e:
            print_error(f"Basic metadata error: {e}")

    def extract_metadata(self, file_path: str) -> None:
        """Dispatches the file to the matching metadata pipeline."""
        console.print(f"\n  [{THEME['PRIMARY']}]→[/] Extracting metadata: "
                      f"[cyan]{os.path.basename(file_path)}[/cyan]")
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            self._extract_pdf_metadata(file_path)
        elif ext in (".jpg", ".jpeg", ".png", ".tiff", ".tif", ".heic", ".webp"):
            self._extract_exif_metadata(file_path)
        elif ext == ".docx":
            self._extract_docx_metadata(file_path)
        elif ext == ".xlsx":
            self._extract_xlsx_metadata(file_path)
        else:
            print_warn(f"No specialised extractor for '{ext}' — showing filesystem info.")
            self._extract_basic_metadata(file_path)

    # ---------------------- Downloading -----------------------------------

    def download_file(self, url: str, extract_metadata: bool = False) -> str | None:
        """
        Streams a binary file to disk. Returns the local path on success.
        """
        try:
            parsed = urlparse(url)
            filename = os.path.basename(parsed.path)
            if not filename:
                filename = hashlib.md5(url.encode()).hexdigest() + ".bin"
                print_warn(f"No filename in URL — using hash: {filename}")

            full_path = os.path.join(self.directory, filename)
            print_info(f"Downloading [green]{filename}[/green] from {url}")

            resp = requests.get(url, stream=True, timeout=15)
            resp.raise_for_status()

            with open(full_path, "wb") as f:
                for chunk in resp.iter_content(chunk_size=8192):
                    f.write(chunk)
            print_success(f"Saved: {full_path}")

            if extract_metadata:
                self.extract_metadata(full_path)
            return full_path

        except requests.exceptions.RequestException as e:
            print_error(f"Network error downloading {url}: {e}")
        except Exception as e:
            print_error(f"Unexpected error downloading {url}: {e}")
        return None

    def download_file_direct(self, url: str, extract_metadata: bool = False) -> str | None:
        """Thin alias for :meth:`download_file`, kept for call sites that
        historically distinguished a "direct" download path."""
        return self.download_file(url, extract_metadata)
